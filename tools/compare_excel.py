"""
Compare generated Excel output against reference cell by cell.
Usage: python tools/compare_excel.py <generated.xlsx> <reference.xlsx> <sheet_name>
"""
import sys
import openpyxl

def normalize(val):
    if val is None:
        return ""
    return str(val).strip()

def compare(gen_path, ref_path, sheet_name):
    ref_wb = openpyxl.load_workbook(ref_path, data_only=True)
    gen_wb = openpyxl.load_workbook(gen_path, data_only=True)

    if sheet_name not in ref_wb.sheetnames:
        print(f"ERROR: sheet '{sheet_name}' not in reference. Available: {ref_wb.sheetnames}")
        return

    if sheet_name not in gen_wb.sheetnames:
        print(f"ERROR: sheet '{sheet_name}' not in generated. Available: {gen_wb.sheetnames}")
        return

    ref_ws = ref_wb[sheet_name]
    gen_ws = gen_wb[sheet_name]

    total = 0
    correct = 0
    wrong = []
    missing = []
    extra = []

    for row in ref_ws.iter_rows():
        for cell in row:
            ref_val = normalize(cell.value)
            if not ref_val:
                continue
            total += 1
            r, c = cell.row, cell.column
            gen_cell = gen_ws.cell(row=r, column=c)
            gen_val = normalize(gen_cell.value)
            if ref_val == gen_val:
                correct += 1
            elif not gen_val:
                missing.append((r, cell.column_letter, ref_val))
            else:
                wrong.append((r, cell.column_letter, ref_val, gen_val))

    # Count extra cells in generated but not in reference
    for row in gen_ws.iter_rows():
        for cell in row:
            gen_val = normalize(cell.value)
            if not gen_val:
                continue
            r, c = cell.row, cell.column
            ref_val = normalize(ref_ws.cell(row=r, column=c).value)
            if not ref_val:
                extra.append((r, cell.column_letter, gen_val))

    pct = correct / total * 100 if total else 0
    print(f"\n=== {sheet_name}: {correct}/{total} = {pct:.1f}% ===")

    if wrong:
        print(f"\nWRONG ({len(wrong)} cells):")
        for r, col, ref, got in wrong[:50]:
            print(f"  R{r}{col}: ref='{ref}'  got='{got}'")
        if len(wrong) > 50:
            print(f"  ... and {len(wrong)-50} more")

    if missing:
        print(f"\nMISSING ({len(missing)} cells):")
        for r, col, ref in missing[:30]:
            print(f"  R{r}{col}: expected='{ref}'")
        if len(missing) > 30:
            print(f"  ... and {len(missing)-30} more")

    if extra:
        print(f"\nEXTRA ({len(extra)} cells):")
        for r, col, val in extra[:30]:
            print(f"  R{r}{col}: '{val}'")
        if len(extra) > 30:
            print(f"  ... and {len(extra)-30} more")

    print(f"\nSUMMARY: correct={correct} wrong={len(wrong)} missing={len(missing)} extra={len(extra)} total={total}")
    print(f"STATUS: {'PASS' if pct >= 90 else 'FAIL'} (target: 90%)")
    return pct, wrong, missing

if __name__ == "__main__":
    if len(sys.argv) < 4:
        print("Usage: python tools/compare_excel.py <generated.xlsx> <reference.xlsx> <sheet_name>")
        sys.exit(1)
    compare(sys.argv[1], sys.argv[2], sys.argv[3])
