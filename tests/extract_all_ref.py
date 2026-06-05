import openpyxl

wb = openpyxl.load_workbook('source/1M SPP SM175 AUTO FUNCTION All CatID.xlsx', data_only=True)

for name in ['MAG03D0424', 'MAG03D0425', 'MAG03D0426', 'MAG03D0427', 'MAG03D0428']:
    if name not in wb.sheetnames:
        print(f"Skipping {name} (not in workbook)")
        continue
    ws = wb[name]
    print(f"\n===== {name} =====")
    for row in ws.iter_rows(min_row=1, max_row=150, max_col=13, values_only=False):
        name_val = row[5].value
        if name_val:
            r = row[0].row
            v_pn = row[7].value or ''
            v_pp = row[8].value or ''
            settings = row[6].value or ''
            relay = row[10].value or ''
            on_delay = row[11].value or ''
            off_delay = row[12].value or ''
            print(f"R{r}: {name_val} | settings={settings} | PN={v_pn} | PP={v_pp} | relay={relay} | ON={on_delay} | OFF={off_delay}")
