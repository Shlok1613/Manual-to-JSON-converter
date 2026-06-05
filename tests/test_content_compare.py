"""Content-based comparison — matches steps by name, not cell position."""
import openpyxl
import sys
import re

def normalize(s):
    """Normalize whitespace and case for comparison."""
    s = str(s).strip().lower()
    s = re.sub(r'\s+', ' ', s)
    return s

def extract_steps(ws, max_row=200):
    """Extract step groups from worksheet. Each group starts at col F (step_name)."""
    steps = []
    r = 10  # First step row for Layout A
    while r <= max_row:
        name = ws.cell(row=r, column=6).value
        if name is None:
            r += 1
            continue
        name = str(name).strip()
        if not name:
            r += 1
            continue

        step = {
            'name': name,
            'settings': [],
            'voltages': [],
            'leds': [],
            'relay': ws.cell(row=r, column=11).value,
            'on_delay': ws.cell(row=r, column=12).value,
            'off_delay': ws.cell(row=r, column=13).value,
        }
        # Collect multi-row data (settings col 7, voltages col 8, LEDs col 10)
        for i in range(4):
            s = ws.cell(row=r+i, column=7).value
            if s: step['settings'].append(str(s).strip())
            v = ws.cell(row=r+i, column=8).value
            if v: step['voltages'].append(str(v).strip())
            l = ws.cell(row=r+i, column=10).value
            if l: step['leds'].append(str(l).strip())

        steps.append(step)
        r += 5  # Each step block = 4 data rows + 1 blank
    return steps

def compare_by_name(ref_path, ref_sheet, gen_path, gen_sheet):
    ref_ws = openpyxl.load_workbook(ref_path, data_only=True)[ref_sheet]
    gen_ws = openpyxl.load_workbook(gen_path, data_only=True)[gen_sheet]

    ref_steps = extract_steps(ref_ws)
    gen_steps = extract_steps(gen_ws)

    print(f'Reference steps: {len(ref_steps)}')
    print(f'Generated steps: {len(gen_steps)}')

    # Match by normalized step name
    gen_by_name = {}
    for s in gen_steps:
        key = normalize(s['name'])
        if key not in gen_by_name:
            gen_by_name[key] = s

    matched = 0
    voltage_match = 0
    led_match = 0
    total_ref = len(ref_steps)

    for ref_s in ref_steps:
        key = normalize(ref_s['name'])
        if key in gen_by_name:
            gen_s = gen_by_name[key]
            matched += 1
            # Check voltages
            ref_v = [normalize(v) for v in ref_s['voltages']]
            gen_v = [normalize(v) for v in gen_s['voltages']]
            if ref_v == gen_v:
                voltage_match += 1
            else:
                print(f'  VOLTAGE DIFF [{ref_s["name"]}]: ref={ref_s["voltages"]} gen={gen_s["voltages"]}')
            # Check LEDs
            ref_l = [normalize(l) for l in ref_s['leds']]
            gen_l = [normalize(l) for l in gen_s['leds']]
            if ref_l == gen_l:
                led_match += 1
        else:
            print(f'  MISSING step: [{ref_s["name"]}]')

    # Extra steps
    ref_names = {normalize(s['name']) for s in ref_steps}
    extra = [s['name'] for s in gen_steps if normalize(s['name']) not in ref_names]
    if extra:
        print(f'\n  EXTRA steps: {extra[:10]}')

    print(f'\n=== CONTENT-BASED COMPARISON ===')
    print(f'Step name match: {matched}/{total_ref} ({matched/total_ref*100:.0f}%)')
    print(f'Voltage match:   {voltage_match}/{matched} ({voltage_match/matched*100:.0f}% of matched)')
    print(f'LED match:       {led_match}/{matched} ({led_match/matched*100:.0f}% of matched)')

if __name__ == '__main__':
    compare_by_name(
        'source/1M SPP SM175 AUTO FUNCTION All CatID.xlsx', 'MAG03D0424',
        'outputs/test_v5/v5test_MAG03D0424.xlsx', 'MAG03D0424',
    )
