"""Compare generated Excel output against reference Excel."""
import openpyxl
import sys


def compare_sheets(ref_path, ref_sheet, gen_path, gen_sheet):
    ref = openpyxl.load_workbook(ref_path, data_only=True)[ref_sheet]
    gen = openpyxl.load_workbook(gen_path, data_only=True)[gen_sheet]

    ref_cells = {}
    for row in ref.iter_rows():
        for cell in row:
            if cell.value is not None:
                ref_cells[(cell.row, cell.column)] = str(cell.value).strip()

    gen_cells = {}
    for row in gen.iter_rows():
        for cell in row:
            if cell.value is not None:
                gen_cells[(cell.row, cell.column)] = str(cell.value).strip()

    total_ref = len(ref_cells)
    found = 0
    wrong = 0
    missing = 0

    for pos, ref_val in ref_cells.items():
        if pos in gen_cells:
            gen_val = gen_cells[pos]
            if ref_val.lower() == gen_val.lower():
                found += 1
            else:
                wrong += 1
                if wrong <= 10:
                    print(f"  WRONG  R{pos[0]}C{pos[1]}: expected '{ref_val}' got '{gen_val}'")
        else:
            missing += 1
            if missing <= 10:
                print(f"  MISSING R{pos[0]}C{pos[1]}: '{ref_val}'")

    extra = len([p for p in gen_cells if p not in ref_cells])

    accuracy = found / total_ref * 100 if total_ref else 0

    print(f"\n=== COMPARISON RESULT ===")
    print(f"Reference cells:  {total_ref}")
    print(f"Correct match:    {found}  ({found/total_ref*100:.1f}%)")
    print(f"Wrong value:      {wrong}")
    print(f"Missing:          {missing}")
    print(f"Extra (invented): {extra}")
    print(f"ACCURACY:         {accuracy:.1f}%")
    print(f"TARGET:           90%")
    print(f"STATUS:           {'PASS' if accuracy >= 90 else 'FAIL'}")
    return accuracy


if __name__ == "__main__":
    import glob
    ref_path = 'source/1M SPP SM175 AUTO FUNCTION All CatID.xlsx'

    # Find generated file for MAG03D0424
    gen_files = glob.glob('outputs/test_pipeline/*MAG03D0424*.xlsx')
    if not gen_files:
        gen_files = glob.glob('outputs/*MAG03D0424*.xlsx')
    if not gen_files:
        print("ERROR: No generated MAG03D0424 file found")
        sys.exit(1)

    # Use the individual file (not All_CatID)
    gen_path = [f for f in gen_files if 'All_CatID' not in f]
    if not gen_path:
        gen_path = gen_files
    gen_path = gen_path[0]

    print(f"Reference: {ref_path}")
    print(f"Generated: {gen_path}")

    # Check available sheets
    gen_wb = openpyxl.load_workbook(gen_path, data_only=True)
    print(f"Generated sheets: {gen_wb.sheetnames}")
    gen_sheet = gen_wb.sheetnames[0]

    accuracy = compare_sheets(ref_path, 'MAG03D0424', gen_path, gen_sheet)
