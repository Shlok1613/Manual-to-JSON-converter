"""Read reference Excel to understand expected output format."""
import openpyxl
import sys

wb = openpyxl.load_workbook('source/1M SPP SM175 AUTO FUNCTION All CatID.xlsx', data_only=True)
print('Sheets:', wb.sheetnames)

for sheet_name in wb.sheetnames[:2]:  # First 2 sheets
    ws = wb[sheet_name]
    print(f'\n=== {sheet_name} (rows 1-50) ===')
    for row in ws.iter_rows(min_row=1, max_row=50, max_col=13, values_only=False):
        vals = []
        for c in row:
            v = c.value
            if v is not None:
                vals.append(f'C{c.column}:{v}')
        if vals:
            sep = ' | '
            print(f'R{row[0].row}: {sep.join(vals)}')
