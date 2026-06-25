"""Generic comparison: compare any variant's output against reference."""
import openpyxl
import sys

variant = sys.argv[1] if len(sys.argv) > 1 else 'MAG03D0424'
ref_path = 'source/1M SPP SM175 AUTO FUNCTION All CatID.xlsx'
gen_path = f'outputs/{variant}.xlsx'

ref_wb = openpyxl.load_workbook(ref_path, data_only=True)
if variant not in ref_wb.sheetnames:
    print(f"ERROR: sheet '{variant}' not found in reference. Available: {ref_wb.sheetnames}")
    sys.exit(1)
ref = ref_wb[variant]

gen_wb = openpyxl.load_workbook(gen_path, data_only=True)
gen_sheet = gen_wb.sheetnames[0]  # First sheet
gen = gen_wb[gen_sheet]

ref_cells = {}
for row in ref.iter_rows():
    for cell in row:
        if cell.value is not None:
            ref_cells[(cell.row, cell.column)] = str(cell.value).strip()

gen_cells = {}
for row in gen.iter_rows():
    for cell in row:
        if cell.value is not None:
            gen_cells[(cell.row, cell.column)] = str(cell.value).strip()

total = len(ref_cells)
found = wrong = missing = 0
wrongs = []
missings = []

for pos, ref_val in ref_cells.items():
    if pos in gen_cells:
        if ref_val.lower() == gen_cells[pos].lower():
            found += 1
        else:
            wrong += 1
            wrongs.append((pos, ref_val, gen_cells[pos]))
    else:
        missing += 1
        missings.append((pos, ref_val))

extra = len([p for p in gen_cells if p not in ref_cells])

print(f"=== {variant} COMPARISON ===")
print(f"Generated: {gen_path} (sheet: {gen_sheet})")
print(f"Reference: {ref_path} (sheet: {variant})\n")

print("=== ALL WRONG VALUES ===")
for (r, c), exp, got in sorted(wrongs):
    print(f"  R{r}C{c}: exp={repr(exp)} got={repr(got)}")

print(f"\n=== ALL MISSING VALUES ===")
for (r, c), val in sorted(missings):
    print(f"  R{r}C{c}: {repr(val)}")

print(f"\n=== EXTRA (INVENTED) CELLS (first 30) ===")
extras = [(pos, gen_cells[pos]) for pos in gen_cells if pos not in ref_cells]
for (r, c), val in sorted(extras)[:30]:
    print(f"  R{r}C{c}: {repr(val)}")

accuracy = found / total * 100 if total else 0
status = "PASS" if accuracy >= 90 else "FAIL"
print(f"\n=== SUMMARY ===")
print(f"Reference: {total}, Correct: {found} ({accuracy:.1f}%), Wrong: {wrong}, Missing: {missing}, Extra: {extra}")
print(f"STATUS: {status} (target: 90%)")
