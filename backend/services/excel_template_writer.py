# backend/services/excel_template_writer.py
from openpyxl import load_workbook
from pathlib import Path
import re
import json
import time

def fill_template_excel(template_path: Path, output_path: Path, ai_data: dict):
    """
    1. Load the Excel template
    2. Duplicate one base sheet (first one)
    3. Fill structured data (Specifications + TestSteps)
    4. Delete all other sheets
    5. Save new Excel file (safe save)
    """

    wb = load_workbook(template_path)
    base_sheet_name = wb.sheetnames[0]
    base_sheet = wb[base_sheet_name]

    # --- Determine product name ---
    product_name = ai_data.get("Product", "New_Product").strip() or "New_Product"

    # --- Duplicate base sheet ---
    if product_name in wb.sheetnames:
        del wb[product_name]
    new_sheet = wb.copy_worksheet(base_sheet)
    new_sheet.title = product_name

    # --- Clear data area (columns F–M) ---
    for row in new_sheet.iter_rows(min_row=6, max_row=new_sheet.max_row, min_col=6, max_col=13):
        for cell in row:
            cell.value = None

    # --- Helper: Flatten nested dicts/lists ---
    def flatten_value(value):
        if isinstance(value, dict):
            return "\n".join([f"{k}: {flatten_value(v)}" for k, v in value.items()])
        elif isinstance(value, list):
            return "\n".join([flatten_value(v) for v in value])
        else:
            return str(value)

    row_index = 6

    # --- Detect and extract structured JSON from raw_output ---
    if "raw_output" in ai_data:
        raw_text = ai_data["raw_output"].strip()
        json_match = re.search(r"```(?:json)?\s*(\{[\s\S]*?\})\s*```", raw_text, re.IGNORECASE)
        if json_match:
            try:
                ai_data = json.loads(json_match.group(1))
            except Exception:
                pass

    # --- Write data based on structure ---
    if "Specifications" in ai_data or "TestSteps" in ai_data:
        specs = ai_data.get("Specifications", {})
        tests = ai_data.get("TestSteps", [])

        # Write Specifications section
        new_sheet.cell(row=row_index, column=6, value="Specifications")
        row_index += 1
        for key, val in specs.items():
            new_sheet.cell(row=row_index, column=6, value=key)
            new_sheet.cell(row=row_index, column=7, value=flatten_value(val))
            row_index += 1

        # Leave space and add TestSteps header
        row_index += 2
        new_sheet.cell(row=row_index, column=6, value="Test Steps")
        row_index += 1

        # Write Test Steps neatly
        headers = ["Action", "Expected Behavior", "Relay Status", "LED Indicator", "Delay", "Condition Type"]
        for i, h in enumerate(headers, start=6):
            new_sheet.cell(row=row_index, column=i, value=h)
        row_index += 1

        for test in tests:
            new_sheet.cell(row=row_index, column=6, value=test.get("Action"))
            new_sheet.cell(row=row_index, column=7, value=test.get("Expected Behavior"))
            new_sheet.cell(row=row_index, column=8, value=test.get("Relay Status"))
            new_sheet.cell(row=row_index, column=9, value=test.get("LED Indicator"))
            new_sheet.cell(row=row_index, column=10, value=test.get("Delay"))
            new_sheet.cell(row=row_index, column=11, value=test.get("Condition Type"))
            row_index += 1

    else:
        # Fallback for simple key–value outputs
        for key, val in ai_data.items():
            if key.lower() == "product":
                continue
            new_sheet.cell(row=row_index, column=6, value=key)
            new_sheet.cell(row=row_index, column=7, value=flatten_value(val))
            row_index += 1

    # --- Delete all other sheets (keep only the new one) ---
    for sheet_name in wb.sheetnames.copy():
        if sheet_name != product_name:
            del wb[sheet_name]

    # --- Safe save (handles locked files) ---
    output_path.parent.mkdir(parents=True, exist_ok=True)
    for attempt in range(3):
        try:
            wb.save(output_path)
            print(f"[Excel Writer] ✅ Saved structured Excel → {output_path.name}")
            break
        except PermissionError:
            timestamp = time.strftime("%H%M%S")
            alt_path = output_path.with_name(f"{output_path.stem}_{timestamp}{output_path.suffix}")
            print(f"[Excel Writer] ⚠️ File in use, saving as {alt_path.name}")
            output_path = alt_path
            continue

    return output_path
