"""End-to-end test: Verify condition generation pipeline."""
import sys
sys.path.insert(0, r"c:\Users\SHLOK\Desktop\Project\Manual-to-JSON-converter\backend")

from services.pdf_extractor import extract_text
from services.block_segmenter import segment_blocks
from services.table_extractor import extract_tables
from services.spec_parser import parse_specifications
from services.condition_generator import generate_test_conditions
from services.excel_writer import generate_excel
from pathlib import Path
import warnings
warnings.filterwarnings('ignore')

pdf_dir = Path(r"c:\Users\SHLOK\Desktop\Project\Manual-to-JSON-converter\backend\uploads")
pdfs = list(pdf_dir.glob("ext_*_FUnctional*.pdf"))
latest_pdf = max(pdfs, key=lambda p: p.stat().st_mtime)
print(f"PDF: {latest_pdf.name}\n")

pages = extract_text(latest_pdf)
full_text = "\n\n".join(pages)
blocks = segment_blocks(full_text)

output_dir = Path(r"c:\Users\SHLOK\Desktop\Project\Manual-to-JSON-converter\backend\outputs")
test_id = "ext_cond_test"

print(f"{'Machine':16s} {'V.Params':>8s} {'T.Params':>8s} {'Conditions':>10s}")
print("-" * 50)

for b in blocks:
    tables = extract_tables(b['text'])
    specs = parse_specifications(b['text'], tables)
    conditions = generate_test_conditions(specs)
    
    vp = len(specs.get('voltage_parameters', {}))
    tp = len(specs.get('timing_parameters', {}))
    
    print(f"{b['machine']:16s} {vp:8d} {tp:8d} {len(conditions):10d}")
    
    filename = generate_excel(test_id, b['machine'], specs, output_dir, test_conditions=conditions)

print(f"\nExcel files:")
for f in sorted(output_dir.glob(f"{test_id}_*.xlsx")):
    print(f"  {f.name} ({f.stat().st_size:,} bytes)")

# Verify SPPR Excel content
from openpyxl import load_workbook
sppr_file = output_dir / f"{test_id}_SPPR.xlsx"
if sppr_file.exists():
    wb = load_workbook(sppr_file)
    print(f"\n=== SPPR Test_Procedures sheet ===")
    ws = wb['Test_Procedures']
    for r in range(1, min(25, ws.max_row + 1)):
        vals = []
        for c in range(1, 14):
            v = ws.cell(r, c).value
            if v is not None:
                vals.append(f"{chr(64+c)}={str(v)[:35]}")
        if vals:
            print(f"  R{r:2d}: {' | '.join(vals)}")
        else:
            print(f"  R{r:2d}: (empty)")
