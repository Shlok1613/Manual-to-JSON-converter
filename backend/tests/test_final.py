"""Final verification: template match + ZIP PDF support."""
import sys, warnings
warnings.filterwarnings('ignore')
sys.path.insert(0, r"c:\Users\SHLOK\Desktop\Project\Manual-to-JSON-converter\backend")

from services.pdf_extractor import extract_text
from services.block_segmenter import segment_blocks
from services.table_extractor import extract_tables
from services.spec_parser import parse_specifications
from services.condition_generator import generate_test_conditions
from services.excel_writer import generate_excel
from pathlib import Path

# ─── TEST 1: FUnctional_Testing PDF (text-based) ───
pdf_dir = Path(r"c:\Users\SHLOK\Desktop\Project\Manual-to-JSON-converter\backend\uploads")
pdfs = list(pdf_dir.glob("ext_*_FUnctional*.pdf"))
if pdfs:
    latest_pdf = max(pdfs, key=lambda p: p.stat().st_mtime)
    print(f"TEST 1: {latest_pdf.name}")

    pages = extract_text(latest_pdf)
    full_text = "\n\n".join(pages)
    blocks = segment_blocks(full_text)

    output_dir = Path(r"c:\Users\SHLOK\Desktop\Project\Manual-to-JSON-converter\backend\outputs")
    test_id = "ext_final_test"

    print(f"  Blocks: {len(blocks)}")
    print(f"  {'Machine':16s} {'Conditions':>10s}")
    print(f"  {'-'*30}")

    for b in blocks:
        tables = extract_tables(b['text'])
        specs = parse_specifications(b['text'], tables)
        conds = generate_test_conditions(specs)
        print(f"  {b['machine']:16s} {len(conds):10d}")
        generate_excel(test_id, b['machine'], specs, output_dir, test_conditions=conds)

    # Verify SPPR Excel content
    from openpyxl import load_workbook
    sppr_file = output_dir / f"{test_id}_SPPR.xlsx"
    if sppr_file.exists():
        wb = load_workbook(sppr_file)
        ws = wb['Test_Procedures']
        print(f"\n  === SPPR Test_Procedures (Template Match Check) ===")
        for r in range(1, min(30, ws.max_row + 1)):
            vals = []
            for c in range(1, 14):
                v = ws.cell(r, c).value
                if v is not None:
                    vals.append(f"{chr(64+c)}={str(v)[:40]}")
            if vals:
                print(f"    R{r:2d}: {' | '.join(vals)}")
            else:
                print(f"    R{r:2d}: (empty)")
else:
    print("TEST 1: No FUnctional PDF found")

# ─── TEST 2: WI.pdf (ZIP-based) ───
print("\n" + "="*60)
wi_pdfs = list(pdf_dir.glob("*WI*.pdf")) + list(pdf_dir.glob("*wi*.pdf"))
# Also check desktop/project root
wi_pdfs += list(Path(r"c:\Users\SHLOK\Desktop\Project\Manual-to-JSON-converter").glob("*WI*.pdf"))
wi_pdfs += list(Path(r"c:\Users\SHLOK\Desktop\Project").glob("*WI*.pdf"))
wi_pdfs += list(Path(r"c:\Users\SHLOK\Desktop").glob("WI.pdf"))

if wi_pdfs:
    wi_pdf = wi_pdfs[0]
    print(f"TEST 2: {wi_pdf.name} ({wi_pdf.stat().st_size:,} bytes)")

    from services.pdf_extractor import is_zip_with_text
    is_zip = is_zip_with_text(wi_pdf)
    print(f"  ZIP detected: {is_zip}")

    pages = extract_text(wi_pdf)
    print(f"  Pages extracted: {len(pages)}")
    if pages:
        full_text = "\n\n".join(pages)
        blocks = segment_blocks(full_text)
        print(f"  Blocks: {len(blocks)}")
        for b in blocks:
            print(f"    {b['machine']}: {len(b['text'])} chars")
else:
    print("TEST 2: No WI.pdf found - will test when uploaded via API")
    print("  ZIP support code is ready in pdf_extractor.py")
